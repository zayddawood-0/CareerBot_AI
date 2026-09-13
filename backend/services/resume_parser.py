"""
Pure text extraction from an uploaded resume file.
No AI here — this module's only job is turning bytes on disk into a plain
text string. Skill/title/experience extraction happens in ai_service.py.
"""
import os

import fitz  # PyMuPDF
import docx  # python-docx


class UnsupportedResumeFormat(Exception):
    pass


def extract_text(file_path: str) -> str:
    """Dispatch to the right extractor based on file extension."""
    ext = os.path.splitext(file_path)[1].lower()

    if ext == ".pdf":
        return _extract_pdf_text(file_path)
    if ext == ".docx":
        return _extract_docx_text(file_path)

    raise UnsupportedResumeFormat(f"Unsupported resume file type: {ext}. Only .pdf and .docx are accepted.")


def _extract_pdf_text(file_path: str) -> str:
    text_chunks = []
    with fitz.open(file_path) as doc:
        for page in doc:
            text_chunks.append(page.get_text())
    return "\n".join(text_chunks).strip()


def _extract_docx_text(file_path: str) -> str:
    document = docx.Document(file_path)
    paragraphs = [p.text for p in document.paragraphs if p.text.strip()]

    # Resumes often put content (skills tables, experience tables) in tables too
    for table in document.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                paragraphs.append(row_text)

    return "\n".join(paragraphs).strip()
